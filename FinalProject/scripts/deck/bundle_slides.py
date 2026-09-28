"""Bundle exported slide PNGs into slides.pdf and slides.pptx (image per slide + speaker notes).

    python bundle_slides.py <exportDir> <outDir>

<exportDir> is the output of export_slides.mjs (slide-NN.png + slides.json).
Needs Pillow and python-pptx.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

WIDE_16_9 = (Emu(12192000), Emu(6858000))


def main(export_dir: Path, out_dir: Path) -> None:
    slides = json.loads((export_dir / "slides.json").read_text(encoding="utf-8"))
    files = [export_dir / s["file"] for s in slides]

    images = [Image.open(f).convert("RGB") for f in files]
    images[0].save(out_dir / "slides.pdf", save_all=True, append_images=images[1:], resolution=144.0)

    prs = Presentation()
    prs.slide_width, prs.slide_height = WIDE_16_9
    blank = prs.slide_layouts[6]
    for slide, file in zip(slides, files, strict=True):
        page = prs.slides.add_slide(blank)
        page.shapes.add_picture(str(file), 0, 0, prs.slide_width, prs.slide_height)
        page.notes_slide.notes_text_frame.text = slide["notes"]
    prs.save(out_dir / "slides.pptx")
    print(f"{len(slides)} slides -> {out_dir / 'slides.pdf'}, {out_dir / 'slides.pptx'}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(Path(sys.argv[1]), Path(sys.argv[2]))
